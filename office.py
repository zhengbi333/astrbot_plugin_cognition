"""Office 文档读写：docx / xlsx。

只做**结构化摘要读**与**从零写**，不做格式精修（那是排版工具的事）。
依赖 python-docx / openpyxl（AstrBot 的 Python 环境自带；缺失时给出可读错误）。
所有路径由调用方（Workspace）先行校验，这里不负责路径安全。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

try:
    from docx import Document as _Document  # type: ignore
except Exception:  # pragma: no cover
    _Document = None

try:
    from openpyxl import Workbook as _Workbook, load_workbook as _load_workbook  # type: ignore
except Exception:  # pragma: no cover
    _Workbook = None
    _load_workbook = None


class OfficeError(Exception):
    """Office 读写失败（信息面向模型可读）。"""


@dataclass(slots=True)
class DocReadResult:
    kind: str
    paragraphs: int
    tables: int
    text: str


def docx_available() -> bool:
    return _Document is not None


def xlsx_available() -> bool:
    return _load_workbook is not None and _Workbook is not None


def _require_docx() -> None:
    if _Document is None:
        raise OfficeError("缺少 python-docx 依赖，无法读写 .docx")


def _require_xlsx() -> None:
    if _load_workbook is None or _Workbook is None:
        raise OfficeError("缺少 openpyxl 依赖，无法读写 .xlsx")


# ---------------------------------------------------------------- docx 读
def read_docx(path: Path, *, max_paragraphs: int = 200, max_chars: int = 8000) -> DocReadResult:
    _require_docx()
    try:
        document = _Document(str(path))
    except Exception as exc:
        raise OfficeError(f"打不开 docx：{type(exc).__name__}: {str(exc)[:120]}") from exc

    lines: list[str] = []
    used = 0
    paragraphs = 0
    for para in document.paragraphs:
        paragraphs += 1
        text = (para.text or "").strip()
        if not text:
            continue
        if len(lines) >= max_paragraphs or used >= max_chars:
            lines.append("…（内容过长，已截断）")
            break
        lines.append(text)
        used += len(text)

    tables = 0
    for table in document.tables:
        tables += 1
        if used >= max_chars:
            break
        rows: list[str] = []
        for row in table.rows[:12]:
            cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            block = "\n".join(rows)
            lines.append(f"[表 {tables}]\n{block}")
            used += len(block)

    body = "\n".join(lines).strip()
    return DocReadResult(kind="docx", paragraphs=paragraphs, tables=tables, text=body)


# ---------------------------------------------------------------- docx 写
def write_docx(path: Path, *, title: str = "", lines: list[str] | None = None) -> int:
    _require_docx()
    document = _Document()
    if title:
        document.add_heading(title, level=1)
    count = 0
    for raw in lines or []:
        text = str(raw).rstrip()
        if not text:
            continue
        if text.startswith("- "):
            document.add_paragraph(text[2:], style="List Bullet")
        else:
            document.add_paragraph(text)
        count += 1
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(path))
    except Exception as exc:
        raise OfficeError(f"保存 docx 失败：{type(exc).__name__}: {str(exc)[:120]}") from exc
    return count


# ---------------------------------------------------------------- xlsx 读写
def read_xlsx(path: Path, *, max_rows: int = 40, max_cols: int = 12, max_chars: int = 8000) -> DocReadResult:
    _require_xlsx()
    try:
        book = _load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise OfficeError(f"打不开 xlsx：{type(exc).__name__}: {str(exc)[:120]}") from exc

    chunks: list[str] = []
    used = 0
    total_rows = 0
    sheet_names = list(getattr(book, "sheetnames", []) or [])
    try:
        for name in sheet_names:
            sheet = book[name]
            rows_out: list[str] = []
            for index, row in enumerate(sheet.iter_rows(max_col=max_cols, values_only=True), 1):
                total_rows += 1
                if index > max_rows or used >= max_chars:
                    rows_out.append("…（该表已截断）")
                    break
                cells = ["" if value is None else str(value) for value in row]
                line = " | ".join(cells).rstrip(" |")
                if line.strip():
                    rows_out.append(line)
                    used += len(line)
            if rows_out:
                chunks.append(f"[工作表 {name}]\n" + "\n".join(rows_out))
            if used >= max_chars:
                break
    finally:
        try:
            book.close()
        except Exception:
            pass
    return DocReadResult(kind="xlsx", paragraphs=total_rows, tables=len(sheet_names), text="\n\n".join(chunks))


def write_xlsx(path: Path, rows: list[list], *, sheet_name: str = "Sheet1") -> int:
    _require_xlsx()
    book = _Workbook()
    sheet = book.active
    if sheet_name:
        try:
            sheet.title = str(sheet_name)[:31] or "Sheet1"
        except Exception:
            pass
    written = 0
    for row in rows or []:
        if not isinstance(row, (list, tuple)):
            row = [row]
        sheet.append([("" if cell is None else cell) for cell in row])
        written += 1
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(str(path))
    except Exception as exc:
        raise OfficeError(f"保存 xlsx 失败：{type(exc).__name__}: {str(exc)[:120]}") from exc
    return written


def parse_rows(raw: str) -> list[list]:
    """把模型给的简单表格文本解析成行列表。

    支持两种写法：
    - 每行用 `|` 分隔（可带首尾竖线）
    - 每行用制表符或连续两个以上空格分隔
    """
    rows: list[list] = []
    for line in str(raw or "").splitlines():
        text = line.rstrip()
        if not text.strip():
            continue
        if "|" in text:
            cells = [c.strip() for c in text.strip().strip("|").split("|")]
        elif "\t" in text:
            cells = [c.strip() for c in text.split("\t")]
        else:
            cells = [c.strip() for c in re.split(r"\s{2,}", text.strip())]
        # 跳过 markdown 表格的分隔行（---|---）
        if cells and all(re.fullmatch(r":?-{2,}:?", c or "") for c in cells):
            continue
        rows.append(cells)
    return rows
